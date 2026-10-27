"""Integration tests for the full compliance validation pipeline."""

import uuid
from unittest.mock import AsyncMock, MagicMock, patch
from docx import Document
from docx.shared import Mm
import pytest
from fastapi.testclient import TestClient

from app.api.dependencies import get_db
from app.api.schemas.report import DocumentStatsSchema, ReportResponse, ViolationSchema
from app.domain.models.job import Job
from app.domain.models.report import Report
from app.main import app
from app.services.profile_service import ProfileService
from app.services.validation_pipeline import ValidationPipeline


@pytest.fixture
def anyio_backend():
    return "asyncio"


def test_pipeline_detects_indent_error(tmp_path):
    """
    Verifies that a .docx with missing paragraph indent is identified as invalid,
    producing ERROR violations and formatted structures ready for DB persistence.
    """
    # 1. Create a dummy docx with invalid indent (0 mm instead of 12.5 mm)
    docx_file = tmp_path / "invalid_indent.docx"
    doc = Document()
    p = doc.add_paragraph("Тестовый абзац без установленного красного отступа.")
    # No first_line_indent set -> indent_mm is 0.0 mm
    doc.save(str(docx_file))

    # 2. Configure rule profile enabling PARAGRAPH_INDENT
    profile_config = {
        "PARAGRAPH_INDENT": {
            "enabled": True,
            "expected_mm": 12.5,
            "tolerance_mm": 1.0,
        }
    }

    # 3. Execute validation pipeline
    pipeline = ValidationPipeline()
    is_valid, total_errors, total_warnings, violations, stats = pipeline.process_file(
        str(docx_file), profile_config
    )

    # 4. Verify output contracts
    assert is_valid is False
    assert total_errors == 1
    assert total_warnings == 0
    assert len(violations) == 1

    violation = violations[0]
    assert violation["rule_id"] == "PARAGRAPH_INDENT"
    assert violation["severity"] == "ERROR"
    assert "Неверный абзацный отступ" in violation["message"]
    assert violation["actual_value"] == "0.0 мм"
    assert violation["expected_value"] == "12.5 мм"
    assert violation["location"]["paragraph_index"] == 0

    # Verify document stats
    assert stats["total_paragraphs"] == 1
    assert stats["total_tables"] == 0
    assert stats["total_sections"] == 1

    # 5. Verify payload can be validated through Pydantic ViolationSchema
    validated_v = ViolationSchema.model_validate(violation)
    assert validated_v.rule_id == "PARAGRAPH_INDENT"

    # 6. Verify payload can be stored in Report ORM model
    test_job_id = uuid.uuid4()
    report_orm = Report(
        id=uuid.uuid4(),
        job_id=test_job_id,
        is_valid=is_valid,
        total_errors=total_errors,
        total_warnings=total_warnings,
        violations=violations,
    )
    assert report_orm.job_id == test_job_id
    assert report_orm.is_valid is False
    assert len(report_orm.violations) == 1

    # 7. Verify serialization into ReportResponse
    report_resp = ReportResponse(
        job_id=str(test_job_id),
        filename="invalid_indent.docx",
        is_valid=report_orm.is_valid,
        summary="Документ содержит ошибки.",
        total_errors=report_orm.total_errors,
        total_warnings=report_orm.total_warnings,
        stats=DocumentStatsSchema(**stats),
        violations=[ViolationSchema.model_validate(v) for v in report_orm.violations],
    )
    assert report_resp.is_valid is False
    assert report_resp.total_errors == 1


def test_pipeline_valid_document(tmp_path):
    """Verifies that a .docx adhering to configured rules returns is_valid == True."""
    docx_file = tmp_path / "valid_document.docx"
    doc = Document()
    p = doc.add_paragraph("Тестовый абзац с корректным отступом 12.5 мм.")
    p.paragraph_format.first_line_indent = Mm(12.5)
    doc.save(str(docx_file))

    profile_config = {
        "PARAGRAPH_INDENT": {
            "enabled": True,
            "expected_mm": 12.5,
            "tolerance_mm": 1.0,
        }
    }

    pipeline = ValidationPipeline()
    is_valid, total_errors, total_warnings, violations, stats = pipeline.process_file(
        str(docx_file), profile_config
    )

    assert is_valid is True
    assert total_errors == 0
    assert total_warnings == 0
    assert len(violations) == 0
    assert stats["total_paragraphs"] == 1


def test_profile_service_load_yaml_preset():
    """Verifies that department_vkr_2025.yaml loads all 16 defined rules."""
    rules = ProfileService.load_yaml_preset()
    assert isinstance(rules, dict)
    assert len(rules) == 16

    expected_rules = [
        "PARAGRAPH_INDENT",
        "PAGE_MARGINS",
        "MAIN_TEXT_FONT",
        "HEADINGS_FORMAT",
        "REQUIRED_SECTIONS",
        "CHAPTERS_COUNT",
        "TOC_MATCHING",
        "FIGURES_FORMAT",
        "TABLES_FORMAT",
        "LISTINGS_FORMAT",
        "APPENDICES_FORMAT",
        "SOURCES_COUNT",
        "CROSS_REFERENCES",
        "LISTS_FORMAT",
        "ANTI_PLAGIARISM_HOMOGLYPHS",
        "FORBIDDEN_ANGLICISMS",
    ]
    for rule_name in expected_rules:
        assert rule_name in rules
        assert rules[rule_name].get("enabled") is True


@pytest.mark.anyio
async def test_profile_service_fallback_when_db_profile_none():
    """Verifies that ProfileService falls back to YAML preset when profile_id is None."""
    config = await ProfileService.get_active_profile_config(session=None, profile_id=None)
    assert "PARAGRAPH_INDENT" in config
    assert config["PARAGRAPH_INDENT"]["expected_mm"] == 12.5


@pytest.mark.anyio
async def test_profile_service_uses_db_profile_when_found():
    """Verifies that ProfileService returns rules_config from DB when found."""
    mock_session = AsyncMock()
    mock_profile = MagicMock()
    mock_profile.rules_config = {
        "PARAGRAPH_INDENT": {"enabled": True, "expected_mm": 15.0}
    }
    mock_result = MagicMock()
    mock_result.scalar_one_or_none.return_value = mock_profile
    mock_session.execute.return_value = mock_result

    config = await ProfileService.get_active_profile_config(
        session=mock_session, profile_id=42
    )
    assert config["PARAGRAPH_INDENT"]["expected_mm"] == 15.0


def test_reports_api_endpoint_scenarios():
    """Verifies the GET /api/v1/reports/{job_id} endpoint responses across statuses."""
    client = TestClient(app)
    mock_db = AsyncMock()

    # Case 1: Job not found -> 404
    async def override_get_db_not_found():
        mock_session = AsyncMock()
        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = None
        mock_session.execute.return_value = mock_result
        yield mock_session

    app.dependency_overrides[get_db] = override_get_db_not_found
    random_job_id = uuid.uuid4()
    resp = client.get(f"/api/v1/reports/{random_job_id}")
    assert resp.status_code == 404

    # Case 2: Job PROCESSING -> 202
    async def override_get_db_processing():
        mock_session = AsyncMock()
        job = Job(id=random_job_id, status="PROCESSING")
        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = job
        mock_session.execute.return_value = mock_result
        yield mock_session

    app.dependency_overrides[get_db] = override_get_db_processing
    resp = client.get(f"/api/v1/reports/{random_job_id}")
    assert resp.status_code == 202
    data = resp.json()
    assert data["status"] == "PROCESSING"

    # Case 3: Job FAILED -> 200 with error info
    async def override_get_db_failed():
        mock_session = AsyncMock()
        job = Job(id=random_job_id, status="FAILED", error_message="File corrupt")
        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = job
        mock_session.execute.return_value = mock_result
        yield mock_session

    app.dependency_overrides[get_db] = override_get_db_failed
    resp = client.get(f"/api/v1/reports/{random_job_id}")
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "FAILED"
    assert data["error"] == "File corrupt"

    # Case 4: Job COMPLETED with Report -> 200 with ReportResponse
    async def override_get_db_completed():
        mock_session = AsyncMock()
        job = Job(id=random_job_id, status="COMPLETED")
        report = Report(
            id=uuid.uuid4(),
            job_id=random_job_id,
            is_valid=False,
            total_errors=1,
            total_warnings=0,
            violations=[
                {
                    "rule_id": "PARAGRAPH_INDENT",
                    "severity": "ERROR",
                    "message": "Неверный отступ",
                    "location": {"paragraph_index": 0},
                }
            ],
        )
        job.report = report
        mock_result = MagicMock()
        mock_result.scalar_one_or_none.return_value = job
        mock_session.execute.return_value = mock_result
        yield mock_session

    app.dependency_overrides[get_db] = override_get_db_completed
    resp = client.get(f"/api/v1/reports/{random_job_id}")
    assert resp.status_code == 200
    report_data = resp.json()
    assert report_data["job_id"] == str(random_job_id)
    assert report_data["is_valid"] is False
    assert report_data["total_errors"] == 1
    assert len(report_data["violations"]) == 1
    assert report_data["violations"][0]["rule_id"] == "PARAGRAPH_INDENT"

    # Clean up overrides
    app.dependency_overrides.clear()


@patch("app.worker.tasks.worker_session_maker")
def test_validate_vkr_worker_task_success(mock_maker, tmp_path):
    """Verifies that validate_vkr task runs end-to-end, updates Job and Report, and cleans up temp file."""
    docx_file = tmp_path / "worker_test.docx"
    doc = Document()
    doc.add_paragraph("Тестовый абзац для проверки работы Celery воркера.")
    doc.save(str(docx_file))

    job_id = str(uuid.uuid4())
    job = Job(id=uuid.UUID(job_id), status="PENDING")

    mock_session = AsyncMock()
    mock_session.add = MagicMock()
    result_job = MagicMock()
    result_job.scalar_one_or_none.return_value = job
    result_report = MagicMock()
    result_report.scalar_one_or_none.return_value = None

    mock_session.execute.side_effect = [result_job, result_report]
    mock_maker.return_value.__aenter__.return_value = mock_session

    from app.worker.tasks import validate_vkr

    res = validate_vkr(job_id, str(docx_file))
    assert res["status"] == "COMPLETED"
    assert job.status == "COMPLETED"
    assert mock_session.add.called
    assert not docx_file.exists()


@patch("app.worker.tasks.worker_session_maker")
def test_validate_vkr_worker_task_failure(mock_maker, tmp_path):
    """Verifies that on unexpected exception, validate_vkr marks job as FAILED and deletes file."""
    bad_file = tmp_path / "corrupt_file.docx"
    bad_file.write_text("corrupted content", encoding="utf-8")

    job_id = str(uuid.uuid4())
    job = Job(id=uuid.UUID(job_id), status="PENDING")

    mock_session = AsyncMock()
    mock_session.add = MagicMock()
    result_job = MagicMock()
    result_job.scalar_one_or_none.return_value = job
    mock_session.execute.return_value = result_job
    mock_maker.return_value.__aenter__.return_value = mock_session

    from app.worker.tasks import validate_vkr

    with pytest.raises(Exception):
        validate_vkr(job_id, str(bad_file))

    assert not bad_file.exists()

