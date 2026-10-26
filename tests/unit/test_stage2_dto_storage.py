import io
import os
import pytest
from fastapi import UploadFile

from app.core.config import settings
from app.domain.dto.document import DocumentDTO, ParagraphDTO, RunDTO, TableDTO
from app.domain.dto.violation import LocationDTO, ViolationDTO
from app.adapters.storage.local import TempFileManager
from app.adapters.extractors.docx.document_parser import DocxExtractor


@pytest.fixture
def anyio_backend():
    return "asyncio"



def test_run_dto_defaults_and_values():
    run = RunDTO(text="Test run", is_bold=True, is_italic=False, font_size_pt=14.0)
    assert run.text == "Test run"
    assert run.is_bold is True
    assert run.is_italic is False
    assert run.font_size_pt == 14.0

    # Test defaults
    run_default = RunDTO(text="Default run")
    assert run_default.is_bold is False
    assert run_default.is_italic is False
    assert run_default.font_size_pt is None


def test_paragraph_dto():
    run1 = RunDTO(text="Hello", is_bold=True)
    run2 = RunDTO(text=" World", is_bold=False)
    paragraph = ParagraphDTO(
        text="Hello World",
        style_name="Normal",
        alignment="JUSTIFY",
        indent_mm=12.5,
        is_heading=False,
        runs=[run1, run2],
    )
    assert paragraph.text == "Hello World"
    assert paragraph.style_name == "Normal"
    assert paragraph.alignment == "JUSTIFY"
    assert paragraph.indent_mm == 12.5
    assert paragraph.is_heading is False
    assert len(paragraph.runs) == 2


def test_table_dto():
    table = TableDTO(
        rows_count=2,
        cols_count=2,
        text_content=[["Col 1", "Col 2"], ["Val 1", "Val 2"]],
    )
    assert table.rows_count == 2
    assert table.cols_count == 2
    assert len(table.text_content) == 2
    assert table.text_content[0] == ["Col 1", "Col 2"]


def test_document_dto():
    doc = DocumentDTO()
    assert doc.paragraphs == []
    assert doc.tables == []
    assert doc.total_pages is None
    assert doc.words_count is None

    p = ParagraphDTO(text="P1", style_name="Normal")
    t = TableDTO(rows_count=1, cols_count=1, text_content=[["Cell"]])
    doc_populated = DocumentDTO(
        paragraphs=[p],
        tables=[t],
        total_pages=10,
        words_count=1500,
    )
    assert len(doc_populated.paragraphs) == 1
    assert len(doc_populated.tables) == 1
    assert doc_populated.total_pages == 10
    assert doc_populated.words_count == 1500


def test_location_dto():
    loc = LocationDTO()
    assert loc.page_number is None
    assert loc.paragraph_index is None
    assert loc.item_title is None
    assert loc.snippet is None

    loc_custom = LocationDTO(
        page_number=5,
        paragraph_index=12,
        item_title="Таблица 1",
        snippet="Текст с ошибкой",
    )
    assert loc_custom.page_number == 5
    assert loc_custom.paragraph_index == 12
    assert loc_custom.item_title == "Таблица 1"
    assert loc_custom.snippet == "Текст с ошибкой"


def test_violation_dto():
    loc = LocationDTO(page_number=3, paragraph_index=7)
    v = ViolationDTO(
        rule_id="RULE_MARGINS_01",
        message="Левое поле меньше 30 мм",
        location=loc,
        severity="ERROR",
        actual_value="25 mm",
        expected_value="30 mm",
    )
    assert v.rule_id == "RULE_MARGINS_01"
    assert v.severity == "ERROR"
    assert v.message == "Левое поле меньше 30 мм"
    assert v.location.page_number == 3
    assert v.actual_value == "25 mm"
    assert v.expected_value == "30 mm"

    # Test default severity and defaults
    v_default = ViolationDTO(
        rule_id="RULE_FONT_01",
        message="Шрифт не Times New Roman",
        location=loc,
    )
    assert v_default.severity == "WARNING"
    assert v_default.actual_value is None
    assert v_default.expected_value is None


@pytest.mark.anyio
async def test_temp_file_manager_lifecycle(tmp_path):
    # Temporarily point TEMP_STORAGE_PATH to tmp_path
    original_path = settings.TEMP_STORAGE_PATH
    settings.TEMP_STORAGE_PATH = str(tmp_path / "vkr_temp")

    try:
        content = b"Fake docx file binary content"
        upload_file = UploadFile(
            file=io.BytesIO(content),
            filename="my_test_work.docx",
        )

        saved_path = None
        async with TempFileManager(upload_file) as file_path:
            saved_path = file_path
            assert os.path.isabs(saved_path)
            assert os.path.exists(saved_path)
            assert "my_test_work.docx" in saved_path

            # Verify contents written
            with open(saved_path, "rb") as f:
                assert f.read() == content

        # Verify file is deleted after exiting context
        assert not os.path.exists(saved_path)
    finally:
        settings.TEMP_STORAGE_PATH = original_path


@pytest.mark.anyio
async def test_temp_file_manager_suppresses_file_not_found(tmp_path):
    original_path = settings.TEMP_STORAGE_PATH
    settings.TEMP_STORAGE_PATH = str(tmp_path / "vkr_temp")

    try:
        upload_file = UploadFile(
            file=io.BytesIO(b"data"),
            filename="premature_delete.docx",
        )

        async with TempFileManager(upload_file) as file_path:
            assert os.path.exists(file_path)
            # Manually delete file prematurely
            os.remove(file_path)
            assert not os.path.exists(file_path)

        # Exiting context should not raise FileNotFoundError
    finally:
        settings.TEMP_STORAGE_PATH = original_path


def test_docx_extractor_empty_document(tmp_path):
    from docx import Document
    docx_file = tmp_path / "empty.docx"
    Document().save(str(docx_file))

    extractor = DocxExtractor()
    result = extractor.extract(str(docx_file))
    assert isinstance(result, DocumentDTO)
    assert result.paragraphs == []
    assert result.tables == []

